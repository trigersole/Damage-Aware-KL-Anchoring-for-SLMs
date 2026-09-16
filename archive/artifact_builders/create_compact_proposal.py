from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path("/Users/dadihemanth15/Documents/Codex/2026-09-04/le")
OUT_DOCX = ROOT / "output" / "documents" / "G4_Vulnerability_Guided_KL_Anchors_Condensed_Proposal.docx"

FONT = "Arial"
BLACK = "000000"
NAVY = "17365D"
PALE_BLUE = "F4F7FB"
LIGHT_GRAY = "D9D9D9"
MID_GRAY = "666666"


def set_font(run, size=9.8, bold=False, italic=False, color=BLACK):
    run.font.name = FONT
    fonts = run._element.get_or_add_rPr().rFonts
    fonts.set(qn("w:ascii"), FONT)
    fonts.set(qn("w:hAnsi"), FONT)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=45, start=70, bottom=45, end=70):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=LIGHT_GRAY, size=5):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), str(size))
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)


def set_repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def set_cell_width(cell, inches):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(inches * 1440)))
    tc_w.set(qn("w:type"), "dxa")


def set_cell_text(cell, text, *, size=7.8, bold=False, color=BLACK, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    run = p.add_run(text)
    set_font(run, size=size, bold=bold, color=color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(cell)


def add_table(doc, headers, rows, widths, font_size=7.8, first_col_center=False):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    set_repeat_header(table.rows[0])
    for index, (header, width) in enumerate(zip(headers, widths)):
        cell = table.rows[0].cells[index]
        set_cell_width(cell, width)
        set_cell_shading(cell, NAVY)
        set_cell_text(cell, header, size=font_size, bold=True, color="FFFFFF", align=WD_ALIGN_PARAGRAPH.CENTER)
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for index, (value, width) in enumerate(zip(values, widths)):
            set_cell_width(cells[index], width)
            if row_index % 2 == 1:
                set_cell_shading(cells[index], PALE_BLUE)
            align = WD_ALIGN_PARAGRAPH.CENTER if first_col_center and index == 0 else WD_ALIGN_PARAGRAPH.LEFT
            set_cell_text(cells[index], str(value), size=font_size, align=align)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_before = Pt(0)
    spacer.paragraph_format.space_after = Pt(0)
    spacer.paragraph_format.line_spacing = Pt(4)
    spacer.paragraph_format.keep_with_next = True
    return table


def add_body(doc, text, *, size=9.8, space_after=2.0, bold_lead=None, italic=False):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.02
    if bold_lead and text.startswith(bold_lead):
        lead = p.add_run(bold_lead)
        set_font(lead, size=size, bold=True)
        rest = p.add_run(text[len(bold_lead):])
        set_font(rest, size=size, italic=italic)
    else:
        run = p.add_run(text)
        set_font(run, size=size, italic=italic)
    return p


def add_heading(doc, text):
    p = doc.add_paragraph(style="Heading 1")
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(1.5)
    run = p.add_run(text)
    set_font(run, size=11.2, bold=True)
    return p


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run("AI6130  |  Team G4  |  ")
    set_font(run, size=7.3, color=MID_GRAY)
    start = OxmlElement("w:fldChar")
    start.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    field_run = paragraph.add_run()
    set_font(field_run, size=7.3, color=MID_GRAY)
    field_run._r.extend([start, instr, end])


def configure_document(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.48)
    section.bottom_margin = Inches(0.47)
    section.left_margin = Inches(0.58)
    section.right_margin = Inches(0.58)
    section.header_distance = Inches(0.2)
    section.footer_distance = Inches(0.22)

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    normal.font.size = Pt(9.8)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.space_after = Pt(2)
    normal.paragraph_format.line_spacing = 1.0

    title = doc.styles["Title"]
    title.font.name = FONT
    title._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    title._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    title.font.size = Pt(16.5)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title_ppr = title._element.get_or_add_pPr()
    border = title_ppr.find(qn("w:pBdr"))
    if border is not None:
        title_ppr.remove(border)

    heading = doc.styles["Heading 1"]
    heading.font.name = FONT
    heading._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    heading._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    heading.font.size = Pt(11.2)
    heading.font.bold = True
    heading.font.color.rgb = RGBColor(0, 0, 0)

    add_page_number(section.footer.paragraphs[0])


def build_document():
    OUT_DOCX.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure_document(doc)
    core = doc.core_properties
    core.title = "Can Vulnerability-Guided KL Anchors Reduce Catastrophic Forgetting in Parameter-Efficient Fine-Tuning?"
    core.subject = "AI6130 Large Language Models condensed project proposal"
    core.author = "Team G4"
    core.keywords = "LoRA, catastrophic forgetting, KL divergence, anchor selection, Qwen2, GSM8K"

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(2.5)
    set_font(title.add_run("Can Vulnerability-Guided KL Anchors Reduce Catastrophic Forgetting in Parameter-Efficient Fine-Tuning?"), size=16.5, bold=True)

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.paragraph_format.space_after = Pt(1)
    set_font(meta.add_run("AI6130 Large Language Models  |  Team G4  |  Professor Luu Anh Tuan  |  15 September 2026"), size=8.2, bold=True)
    members = doc.add_paragraph()
    members.alignment = WD_ALIGN_PARAGRAPH.CENTER
    members.paragraph_format.space_after = Pt(3)
    set_font(members.add_run("Sridhar Srihari (G2510717L)  |  Siddharth Paliwal (G2510801L)  |  Dadi Hemanth (G2510794D)  |  Balasubramanian Hariharan (G2610112K)"), size=7.8)

    add_heading(doc, "1 Goals Objectives and Originality")
    add_body(
        doc,
        "Scientific question. Under equal target-training, anchor-token, and teacher-compute budgets, can a short pilot LoRA identify unlabeled prompts whose use as KL anchors reduces catastrophic forgetting more effectively than random or diversity-only anchors at matched GSM8K performance? We will adapt Qwen2-1.5B-Instruct to mathematical word problems while preserving unrelated abilities. This stability-plasticity problem matters because inexpensive task adaptation should not silently damage useful behavior. We hypothesize that Damage-KL and Diverse-Damage-KL will retain more than Random-KL, and that diversity will help when high-damage prompts are redundant. LoRA, distillation, KL regularization, and diversity sampling are established; our attempted originality is the controlled combination of pilot-measured behavioral vulnerability, unlabeled anchor selection, fixed protection budgets, target-matched comparisons, and recovery-aware evaluation. A null result will still answer the research question.",
        bold_lead="Scientific question.",
    )

    add_heading(doc, "2 Tasks")
    add_table(
        doc,
        ["Task", "Setup and example", "Output or measure"],
        [
            ("Target adaptation", "Supervised LoRA on GSM8K. Example input: A shop has 18 notebooks and sells 7. How many remain?", "Expected output: 18 - 7 = 11. Final answer: 11."),
            ("Retention", "BoolQ yes-no comprehension; HellaSwag commonsense completion; ARC-Easy science multiple choice.", "Fixed answer-option accuracy before and after adaptation."),
            ("Anchor selection", "About 2,000 unlabeled general instructions. Score teacher-to-pilot response-token KL, then select random, semantic-diverse, high-damage, or diverse-damage anchors.", "Equal anchor-token sets for controlled main training."),
        ],
        [1.15, 4.15, 2.02],
        font_size=8.1,
    )

    add_heading(doc, "3 Data and Preprocessing")
    add_table(
        doc,
        ["Dataset", "Role", "Split and preprocessing"],
        [
            ("GSM8K", "Target", "Train on official training data with a fixed development split; keep the official test untouched; format prompts and extract the final numeric answer."),
            ("Dolly 15k", "Anchor pool", "Use about 2,000 instruction fields without labels; remove empty, unsafe, long, mathematics-like, and duplicate prompts."),
            ("BoolQ", "Retention", "Use held-out training items for protocol checks and official validation for final accuracy."),
            ("HellaSwag", "Retention", "Use official validation with length-normalized option likelihood."),
            ("ARC-Easy", "Retention", "Use official validation with answer-option likelihood."),
        ],
        [1.05, 1.05, 5.22],
        font_size=7.9,
    )
    add_body(doc, "All datasets are public; no new data will be collected. We will record versions and indices and use exact, normalized-text, and embedding-similarity checks to keep training, anchors, development data, and final evaluation disjoint. Benchmark questions and answer choices will never be used as anchors.", size=9.1, space_after=0)

    doc.add_page_break()

    add_heading(doc, "4 Models and Methods")
    add_body(
        doc,
        "Model and procedure. We will use Qwen/Qwen2-1.5B-Instruct with rank-16 LoRA, a frozen backbone, fixed tokenizer, chat template, target modules, sequence length, and optimizer schedule. BF16 is preferred; four-bit loading is a documented, identical fallback. An unchanged copy is the teacher. The student minimizes GSM8K task loss plus a development-selected weight times teacher-to-student KL on anchors. First, evaluate the base and cache deterministic teacher continuations. Second, train a temporary pilot for 5-10 percent of target steps and score average response-token teacher-to-pilot KL. Third, form equal-budget random, diversity, damage, and diverse-damage sets. Fourth, discard the pilot and start every main run from the same base. Fifth, match target examples, anchor tokens, teacher passes, and optimizer work; save checkpoints and evaluate learning, retention, recovery, and cost. The score measures behavioral vulnerability, not damaged parameters or proven knowledge erasure.",
        bold_lead="Model and procedure.",
    )
    add_body(doc, "Downloaded. Qwen2 checkpoint, GSM8K, Dolly 15k, BoolQ, HellaSwag, ARC-Easy, Transformers, Datasets, PEFT, a sentence-embedding model, and the LM Evaluation Harness. Implemented by us. Pilot protocol, continuation cache, token-masked KL loss, four selectors, budget controller, target-matched comparison, recovery tests, item transitions, contamination checks, experiment registry, and analysis.", size=9.1, bold_lead="Downloaded.")

    add_heading(doc, "5 Baselines")
    add_body(doc, "All six conditions will be implemented and run by our team from the same downloaded Qwen2 checkpoint.", size=8.9, space_after=1.5)
    add_table(
        doc,
        ["Condition", "Anchor rule", "Purpose"],
        [
            ("Untouched base", "None", "Initial capability reference."),
            ("Task-only LoRA", "None", "Ordinary learning and forgetting."),
            ("Random-KL", "Uniform random", "Tests whether any anchors suffice."),
            ("Diversity-KL", "Across semantic clusters", "Tests coverage without vulnerability."),
            ("Damage-KL", "Highest pilot KL", "Tests vulnerability alone."),
            ("Diverse-Damage-KL", "High KL across clusters", "Primary proposed selector with redundancy control."),
        ],
        [1.48, 1.72, 4.12],
        font_size=7.9,
    )

    add_heading(doc, "6 Evaluation")
    add_table(
        doc,
        ["Question", "Automatic metric and comparison"],
        [
            ("Target learning", "GSM8K final-answer exact match and negative log-likelihood; report post-minus-base gain."),
            ("Forgetting", "Base-minus-adapted accuracy on BoolQ, HellaSwag, and ARC-Easy; report each task, macro average, worst task, and correct-to-incorrect versus incorrect-to-correct transitions."),
            ("Fairness and efficiency", "Compare checkpoints within about one GSM8K percentage point or the full gain-forgetting frontier; fix anchor tokens and teacher work; report retention gain per 100,000 anchor tokens and GPU-hour."),
            ("Robustness", "At least two seeds, seed-level results, paired-bootstrap 95 percent confidence intervals, alternative-template and fixed few-shot recovery, plus qualitative error categories."),
        ],
        [1.55, 5.77],
        font_size=7.9,
    )
    add_body(doc, "Published comparator. In a different Llama-2-7B/MetaMathQA setting, Biderman et al. reported GSM8K accuracy of 62.2 percent for LoRA and 64.2 percent for full fine-tuning, with retention averages of 0.63 and 0.57 [2]. These are contextual scores; conclusions will use within-Qwen2 baselines under identical settings.", size=8.8, bold_lead="Published comparator.")

    add_heading(doc, "7 Feasibility Risks and Outputs")
    add_body(doc, "The minimum study is five trained LoRA conditions with two seeds, or ten central runs, plus the frozen base and short pilot on a 1.5B model. The middle anchor budget will initially be 128 prompts. We will cache teacher outputs and, if required, reduce candidate size or sequence length or use the same documented top-k approximation for all KL runs. If forgetting is weak, we will extend target exposure or public mathematics-instruction data while keeping GSM8K fixed; if pilot rankings are unstable or KL blocks learning, we will report this and use target matching rather than tune on final tests. Core runs take priority over sweeps. Outputs include reproducible code and manifests, selected anchors, learning-retention Pareto plots, uncertainty, recovery and error analyses, and positive, null, or negative findings. Only after the primary matrix is complete will we attempt a second adaptation dataset or domain.", size=9.1)

    add_heading(doc, "References")
    references = (
        "[1] Hu et al., LoRA, ICLR 2022.  [2] Biderman et al., LoRA Learns Less and Forgets Less, TMLR 2024.  "
        "[3] Hinton et al., Distilling the Knowledge in a Neural Network, 2015.  [4] Li and Hoiem, Learning without Forgetting, ECCV 2016.  "
        "[5] Jin and Ren, Demystifying Language Model Forgetting with Low-rank Example Associations, NeurIPS 2025.  "
        "[6] Zheng et al., Spurious Forgetting in Continual Learning of Language Models, ICLR 2025.  "
        "[7] Yang et al., Qwen2 Technical Report, 2024.  [8] Cobbe et al., Training Verifiers to Solve Math Word Problems, 2021.  "
        "[9] Clark et al., BoolQ, NAACL 2019.  [10] Zellers et al., HellaSwag, ACL 2019.  [11] Clark et al., ARC, 2018.  "
        "[12] Reimers and Gurevych, Sentence-BERT, EMNLP 2019.  [13] Databricks Dolly 15k Dataset Card, 2023."
    )
    ref_paragraph = add_body(doc, references, size=7.4, space_after=0)
    ref_paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT

    doc.save(OUT_DOCX)
    print(OUT_DOCX)


if __name__ == "__main__":
    build_document()
