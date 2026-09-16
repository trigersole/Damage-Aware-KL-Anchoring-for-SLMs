from __future__ import annotations

from pathlib import Path
from textwrap import fill

import numpy as np
from PIL import Image, ImageDraw, ImageFont

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "output" / "documents" / "damage-aware-kl-anchoring-complete-project-guide.docx"
ASSET_DIR = ROOT / "work" / "docx-assets"
WORKFLOW_PNG = ASSET_DIR / "damage-aware-workflow.png"
SELECTION_PNG = ASSET_DIR / "anchor-selection-concept.png"

FONT = "Arial"
MONO = "Courier New"
BLACK = "000000"
NAVY = "1F4E78"
BLUE = "2F75B5"
MID_BLUE = "5B9BD5"
PALE_BLUE = "EAF2F8"
PALE_GRAY = "F5F7F9"
MID_GRAY = "66727D"
GRID = "CDD6DF"
WHITE = "FFFFFF"


def set_run_font(run, name: str = FONT, size: float | None = None,
                 bold: bool | None = None, italic: bool | None = None,
                 color: str | None = None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    return run


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def shade_cell(cell, color: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), color)


def set_table_borders(table, color=GRID, size="5"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), size)
        tag.set(qn("w:space"), "0")
        tag.set(qn("w:color"), color)


def remove_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "nil")


def add_hyperlink(paragraph, text: str, url: str, color=BLUE, underline=True):
    part = paragraph.part
    rel_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    new_run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    r_fonts = OxmlElement("w:rFonts")
    r_fonts.set(qn("w:ascii"), FONT)
    r_fonts.set(qn("w:hAnsi"), FONT)
    r_pr.append(r_fonts)
    c = OxmlElement("w:color")
    c.set(qn("w:val"), color)
    r_pr.append(c)
    if underline:
        u = OxmlElement("w:u")
        u.set(qn("w:val"), "single")
        r_pr.append(u)
    new_run.append(r_pr)
    t = OxmlElement("w:t")
    t.text = text
    new_run.append(t)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)
    return hyperlink


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr_text, fld_char2])
    set_run_font(run, size=8.5, color=MID_GRAY)


def set_alt_text(inline_shape, description: str):
    try:
        inline_shape._inline.docPr.set("descr", description)
    except Exception:
        pass


def create_styles(doc: Document):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    normal.font.size = Pt(10.6)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.space_after = Pt(5.5)
    normal.paragraph_format.line_spacing = 1.12

    title = styles["Title"]
    title.font.name = FONT
    title._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    title._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    title.font.size = Pt(30)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title.paragraph_format.space_after = Pt(15)
    title_ppr = title._element.get_or_add_pPr()
    title_borders = title_ppr.find(qn("w:pBdr"))
    if title_borders is not None:
        title_ppr.remove(title_borders)

    for name, size, before, after in (
        ("Heading 1", 18, 18, 7),
        ("Heading 2", 13.5, 13, 5),
        ("Heading 3", 11.3, 9, 3),
    ):
        style = styles[name]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:ascii"), FONT)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    if "Subtitle" not in styles:
        styles.add_style("Subtitle", WD_STYLE_TYPE.PARAGRAPH)
    subtitle = styles["Subtitle"]
    subtitle.font.name = FONT
    subtitle._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    subtitle._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    subtitle.font.size = Pt(15)
    subtitle.font.color.rgb = RGBColor.from_string(MID_GRAY)
    subtitle.paragraph_format.space_after = Pt(10)

    for custom_name, size, color, italic in (
        ("Lead", 12.2, BLACK, False),
        ("Caption Custom", 8.8, MID_GRAY, True),
        ("Small Note", 9.0, MID_GRAY, False),
        ("Code Block", 8.8, BLACK, False),
    ):
        if custom_name not in styles:
            styles.add_style(custom_name, WD_STYLE_TYPE.PARAGRAPH)
        style = styles[custom_name]
        style.font.name = MONO if custom_name == "Code Block" else FONT
        style._element.rPr.rFonts.set(qn("w:ascii"), style.font.name)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), style.font.name)
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.italic = italic
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.line_spacing = 1.05 if custom_name == "Code Block" else 1.1
        if custom_name == "Code Block":
            style.paragraph_format.left_indent = Inches(0.25)
            style.paragraph_format.right_indent = Inches(0.15)
            style.paragraph_format.keep_together = True

    for list_name in ("List Bullet", "List Number"):
        style = styles[list_name]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:ascii"), FONT)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
        style.font.size = Pt(10.4)
        style.paragraph_format.space_after = Pt(2.7)
        style.paragraph_format.line_spacing = 1.08


def add_body(doc: Document, text: str, style: str | None = None,
             bold_lead: str | None = None, keep=False):
    p = doc.add_paragraph(style=style)
    if bold_lead and text.startswith(bold_lead):
        set_run_font(p.add_run(bold_lead), bold=True)
        set_run_font(p.add_run(text[len(bold_lead):]))
    else:
        set_run_font(p.add_run(text))
    if keep:
        p.paragraph_format.keep_together = True
    return p


def add_bullets(doc: Document, items: list[str], level=0):
    for item in items:
        p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
        set_run_font(p.add_run(item))
    return None


def add_numbers(doc: Document, items: list[str]):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        set_run_font(p.add_run(item))


def add_table(doc: Document, headers: list[str], rows: list[list[str]],
              widths: list[float] | None = None, font_size=8.9):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False if widths else True
    table.style = "Table Grid"
    set_table_borders(table)
    header_row = table.rows[0]
    set_repeat_table_header(header_row)
    for idx, header in enumerate(headers):
        cell = header_row.cells[idx]
        shade_cell(cell, NAVY)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(cell)
        if widths:
            cell.width = Inches(widths[idx])
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_after = Pt(0)
        set_run_font(p.add_run(header), size=9.0, bold=True, color=WHITE)
    for ridx, row_data in enumerate(rows):
        row = table.add_row()
        prevent_row_split(row)
        for cidx, value in enumerate(row_data):
            cell = row.cells[cidx]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            set_cell_margins(cell)
            if widths:
                cell.width = Inches(widths[cidx])
            if ridx % 2 == 1:
                shade_cell(cell, PALE_GRAY)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.04
            set_run_font(p.add_run(str(value)), size=font_size)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_figure(doc: Document, path: Path, width: float, caption: str, alt: str):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    shape = p.add_run().add_picture(str(path), width=Inches(width))
    set_alt_text(shape, alt)
    cp = doc.add_paragraph(style="Caption Custom")
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(cp.add_run(caption), size=8.8, italic=True, color=MID_GRAY)


def part_heading(doc: Document, text: str):
    p = doc.add_paragraph(text, style="Heading 1")
    p.paragraph_format.page_break_before = True
    return p


def add_reference(doc: Document, number: int, title: str, venue: str, url: str, note: str):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.22)
    p.paragraph_format.first_line_indent = Inches(-0.22)
    p.paragraph_format.space_after = Pt(5)
    set_run_font(p.add_run(f"[{number}] "), bold=True)
    add_hyperlink(p, title, url)
    set_run_font(p.add_run(f". {venue}. {note}"), size=9.5)


def configure_page(doc: Document):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.68)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)
    section.header_distance = Inches(0.28)
    section.footer_distance = Inches(0.3)
    section.different_first_page_header_footer = True

    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_after = Pt(0)
    set_run_font(p.add_run("Damage Aware KL Anchoring"), size=8.2, color=BLACK)

    footer = section.footer
    table = footer.add_table(rows=1, cols=2, width=Inches(6.85))
    remove_table_borders(table)
    table.columns[0].width = Inches(5.8)
    table.columns[1].width = Inches(1.05)
    left = table.cell(0, 0).paragraphs[0]
    left.paragraph_format.space_after = Pt(0)
    set_run_font(left.add_run("Team project working document"), size=8.2, color=MID_GRAY)
    add_page_number(table.cell(0, 1).paragraphs[0])


def make_workflow_figure():
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    width, height = 1600, 2080
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    regular_path = "/System/Library/Fonts/Supplemental/Arial.ttf"
    bold_path = "/System/Library/Fonts/Supplemental/Arial Bold.ttf"
    title_font = ImageFont.truetype(bold_path, 49)
    stage_font = ImageFont.truetype(bold_path, 34)
    body_font = ImageFont.truetype(regular_path, 27)
    stages = [
        ("1  Establish the reference", "Freeze the base model and record target and retention performance"),
        ("2  Run a short pilot", "Train a temporary LoRA adapter on a small math subset"),
        ("3  Measure vulnerability", "Score 2,000 general prompts by base to pilot output drift"),
        ("4  Build equal size anchor sets", "Random  |  Diversity  |  Damage  |  Diverse damage"),
        ("5  Discard the pilot and reset", "Start every main condition from the same untouched base model"),
        ("6  Train the main students", "Math loss plus KL protection on the selected anchors"),
        ("7  Compare at matched learning", "Math gain  |  Retention  |  Recovery  |  Cost"),
    ]
    heading = "The complete experimental workflow"
    bbox = draw.textbbox((0, 0), heading, font=title_font)
    draw.text(((width - (bbox[2] - bbox[0])) / 2, 35), heading, font=title_font, fill="#000000")
    box_left, box_right = 150, 1450
    box_h, gap = 185, 91
    top = 130
    for idx, (title, subtitle) in enumerate(stages):
        y1 = top + idx * (box_h + gap)
        y2 = y1 + box_h
        draw.rounded_rectangle((box_left, y1, box_right, y2), radius=18,
                               fill="white", outline=f"#{NAVY}", width=5)
        draw.text((205, y1 + 32), title, font=stage_font, fill="#000000")
        lines = fill(subtitle, 74).splitlines()
        for line_idx, line in enumerate(lines):
            draw.text((205, y1 + 91 + line_idx * 31), line, font=body_font, fill=f"#{MID_GRAY}")
        if idx < len(stages) - 1:
            x = width // 2
            start_y = y2 + 10
            end_y = y2 + gap - 12
            draw.line((x, start_y, x, end_y - 20), fill=f"#{BLUE}", width=5)
            draw.polygon([(x - 14, end_y - 24), (x + 14, end_y - 24), (x, end_y)], fill=f"#{BLUE}")
    image.save(WORKFLOW_PNG, dpi=(220, 220))


def make_selection_figure():
    rng = np.random.default_rng(7)
    centers = np.array([[-2.0, 1.2], [1.8, 1.8], [-1.2, -1.5], [2.1, -1.1]])
    points = []
    clusters = []
    for idx, center in enumerate(centers):
        pts = center + rng.normal(scale=0.48, size=(24, 2))
        points.append(pts)
        clusters.extend([idx] * len(pts))
    pts = np.vstack(points)
    clusters = np.array(clusters)
    damage = np.clip(rng.beta(1.6, 3.2, size=len(pts)), 0.02, 0.98)
    damage[[7, 34, 57, 83]] = [0.96, 0.91, 0.88, 0.94]

    selections = {
        "Random": rng.choice(len(pts), size=8, replace=False),
        "Diversity": np.concatenate([rng.choice(np.where(clusters == c)[0], 2, replace=False) for c in range(4)]),
        "Damage": np.argsort(damage)[-8:],
        "Diverse damage": np.concatenate([
            np.where(clusters == c)[0][np.argsort(damage[clusters == c])[-2:]] for c in range(4)
        ]),
    }

    width, height = 1800, 1480
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    regular_path = "/System/Library/Fonts/Supplemental/Arial.ttf"
    bold_path = "/System/Library/Fonts/Supplemental/Arial Bold.ttf"
    title_font = ImageFont.truetype(bold_path, 41)
    panel_font = ImageFont.truetype(bold_path, 31)
    note_font = ImageFont.truetype(regular_path, 24)
    palette = ["#7EA6C9", "#7FB7A6", "#C6A36A", "#A991C7"]
    heading = "How four selectors can spend the same eight prompt budget"
    bbox = draw.textbbox((0, 0), heading, font=title_font)
    draw.text(((width - (bbox[2] - bbox[0])) / 2, 25), heading, font=title_font, fill="#000000")
    panel_boxes = [(80, 120, 860, 720), (940, 120, 1720, 720),
                   (80, 775, 860, 1375), (940, 775, 1720, 1375)]
    min_x, max_x = -3.2, 3.2
    min_y, max_y = -2.8, 3.0
    selected_set_lookup = {name: set(int(x) for x in values) for name, values in selections.items()}
    for (title, selected), (x1, y1, x2, y2) in zip(selections.items(), panel_boxes):
        draw.rectangle((x1, y1, x2, y2), fill="white", outline=f"#{GRID}", width=4)
        tb = draw.textbbox((0, 0), title, font=panel_font)
        draw.text(((x1 + x2 - (tb[2] - tb[0])) / 2, y1 + 18), title, font=panel_font, fill="#000000")
        plot_left, plot_right = x1 + 45, x2 - 45
        plot_top, plot_bottom = y1 + 78, y2 - 34
        chosen = selected_set_lookup[title]
        for idx, ((px, py), c, d) in enumerate(zip(pts, clusters, damage)):
            sx = plot_left + (px - min_x) / (max_x - min_x) * (plot_right - plot_left)
            sy = plot_bottom - (py - min_y) / (max_y - min_y) * (plot_bottom - plot_top)
            radius = int(7 + d * 15)
            draw.ellipse((sx - radius, sy - radius, sx + radius, sy + radius),
                         fill=palette[int(c)], outline=None)
            if idx in chosen:
                ring = radius + 8
                draw.ellipse((sx - ring, sy - ring, sx + ring, sy + ring),
                             fill=None, outline="#C00000", width=5)
    note = "Nearby dots have similar meaning  •  Larger dots have higher pilot damage  •  Red circles are selected anchors"
    nb = draw.textbbox((0, 0), note, font=note_font)
    draw.text(((width - (nb[2] - nb[0])) / 2, 1414), note, font=note_font, fill=f"#{MID_GRAY}")
    image.save(SELECTION_PNG, dpi=(220, 220))


def add_cover(doc: Document):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(54)
    p.paragraph_format.space_after = Pt(18)
    set_run_font(p.add_run("RESEARCH ORIENTED LLM COURSE PROJECT"), size=10.2, bold=True, color=MID_GRAY)

    p = doc.add_paragraph(style="Title")
    set_run_font(p.add_run("Damage Aware KL Anchoring for Small Language Models"), size=30, bold=True, color=BLACK)

    p = doc.add_paragraph(style="Subtitle")
    set_run_font(p.add_run("Complete project explanation experimental design and team plan"), size=15, color=MID_GRAY)

    doc.add_paragraph().paragraph_format.space_after = Pt(12)
    p = doc.add_paragraph(style="Lead")
    p.paragraph_format.right_indent = Inches(0.65)
    set_run_font(p.add_run(
        "We will test whether a short pilot fine tuning run can identify general prompts that are especially vulnerable to a mathematics LoRA update, and whether protecting those prompts with a frozen teacher preserves more capability than random or diversity based anchoring under the same training budget."
    ), size=12.2)

    doc.add_paragraph().paragraph_format.space_after = Pt(18)
    meta = doc.add_table(rows=4, cols=2)
    meta.alignment = WD_TABLE_ALIGNMENT.LEFT
    meta.autofit = False
    remove_table_borders(meta)
    labels = ["Audience", "Core scope", "Team and duration", "Status"]
    values = [
        "MSAI teammates and course instructor",
        "One small model  •  GSM8K adaptation  •  Five LoRA conditions  •  Two seeds",
        "Three to four students  •  Eight weeks",
        "Research proposal and implementation guide",
    ]
    for i, (label, value) in enumerate(zip(labels, values)):
        meta.cell(i, 0).width = Inches(1.35)
        meta.cell(i, 1).width = Inches(5.25)
        for cell in meta.rows[i].cells:
            set_cell_margins(cell, top=70, bottom=70, start=0, end=100)
        p1 = meta.cell(i, 0).paragraphs[0]
        set_run_font(p1.add_run(label), size=9.5, bold=True)
        p2 = meta.cell(i, 1).paragraphs[0]
        set_run_font(p2.add_run(value), size=9.5, color=MID_GRAY)

    doc.add_paragraph().paragraph_format.space_after = Pt(34)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(30)
    p.paragraph_format.keep_together = True
    set_run_font(p.add_run("Recommended project title"), size=9.5, bold=True, color=MID_GRAY)
    p.add_run().add_break()
    set_run_font(p.add_run("Which Prompts Should a Fine Tuned Model Remember"), size=15, bold=True, color=BLACK)
    p.add_run().add_break()
    set_run_font(p.add_run("Damage aware KL anchoring with recovery aware evaluation"), size=10.5, color=MID_GRAY)

    doc.add_page_break()


def build_document():
    make_workflow_figure()
    make_selection_figure()

    doc = Document()
    create_styles(doc)
    configure_page(doc)

    props = doc.core_properties
    props.title = "Damage Aware KL Anchoring for Small Language Models"
    props.subject = "A complete research proposal and implementation guide for an LLM course project"
    props.author = "MSAI course project team"
    props.keywords = "LoRA, catastrophic forgetting, KL divergence, knowledge distillation, anchor selection"

    add_cover(doc)

    doc.add_heading("Project decision", level=1)
    add_body(doc,
        "We recommend a focused method and evaluation project called damage aware KL anchoring. We adapt a small language model to mathematical word problems with LoRA. A frozen copy of the original model acts as a teacher. During training, a limited set of general prompts reminds the student to preserve its earlier behaviour. The research contribution is how those prompts are selected and how forgetting is measured.",
        style="Lead")
    add_body(doc,
        "The ordinary baseline selects protective prompts randomly. Our method first performs a short temporary mathematics fine tuning run. It then measures which general prompts changed most between the original model and the temporary model. Those vulnerable prompts become candidates for protection in the real training run. The pilot model is discarded before the main experiment.")
    add_body(doc,
        "The primary question is whether pilot guided anchors preserve more general capability than random anchors when both methods receive the same number of anchor tokens and reach similar mathematics performance. We also test whether the apparent loss can be recovered by changing the prompt or answer scoring method. This prevents us from calling every benchmark drop permanent knowledge erasure.")

    doc.add_heading("Research question", level=2)
    add_body(doc,
        "Given a fixed number of unlabeled anchor tokens and equal target task training, does selecting prompts by early base to adapted output drift preserve more general capability than random or diversity only KL anchoring, and does any advantage remain after recovery controls?",
        keep=True)

    doc.add_heading("Core recommendation", level=2)
    add_table(doc,
        ["Decision", "Recommended choice", "Reason"],
        [
            ["Primary model", "Qwen2.5 0.5B Instruct", "Small enough for repeated runs while still supporting instruction and reasoning evaluations"],
            ["Target task", "GSM8K mathematical word problems", "Public data, clear numeric scoring, and a narrow adaptation objective"],
            ["Central conditions", "Task only, Random KL, Diversity KL, Damage KL, Diverse Damage KL", "One baseline, two controls, and two proposed selectors"],
            ["Repetitions", "Two seeds per condition", "Ten central runs are realistic for an eight week student project"],
            ["Primary comparison", "Diverse Damage KL versus Random KL", "Tests whether vulnerability plus coverage beats ordinary anchoring"],
            ["Primary fairness rule", "Compare checkpoints at similar GSM8K validation performance", "A method must not look safer merely because it learned less mathematics"],
        ], widths=[1.35, 2.05, 3.65], font_size=8.7)

    doc.add_heading("What makes the idea research oriented", level=2)
    add_body(doc,
        "LoRA, knowledge distillation, KL regularization, replay, and targeted example selection already exist. We should not claim that any one of these ingredients is new. Our research question sits at their controlled intersection: use an early behavioural damage signal to select a small unlabeled KL anchor set, compare it with random and meaning based coverage under equal budgets, match target learning, and separate recoverable failures from residual loss.")
    add_body(doc,
        "A bounded literature review completed for this project screened 180 records and examined 60 papers in greater detail. It found close precedents for damage aware labeled rehearsal, example level forgetting prediction, and KL approximations, but no direct controlled match for this exact small model protocol. Novelty confidence is therefore medium, not certain. We should repeat a focused search immediately before submitting the proposal and describe the contribution as an extension and controlled study rather than a first ever method.")

    doc.add_heading("Guide map", level=2)
    add_table(doc,
        ["Part", "What it explains"],
        [
            ["A Foundations", "AI, models, parameters, tokens, probability distributions, fine tuning, and LoRA"],
            ["B Problem and protection", "Capability degradation, the frozen teacher, anchor prompts, and KL divergence"],
            ["C Proposed method", "The pilot, damage scores, four selection strategies, reset, and main training"],
            ["D Experimental design", "Conditions, data, budgets, target matching, recovery evaluation, and statistics"],
            ["E Implementation", "Training logic, cached teacher outputs, files, checks, and reproducibility"],
            ["F Execution", "Eight week plan, team ownership, risks, fallbacks, and decision rules"],
            ["G Communication", "Expected interpretations, questions from the professor, glossary, abstract, and references"],
        ], widths=[1.45, 5.6], font_size=8.9)

    part_heading(doc, "Part A Foundations")
    doc.add_heading("Artificial intelligence and machine learning", level=2)
    add_body(doc,
        "Artificial intelligence is the broad field of building computer systems that perform tasks associated with human intelligence, such as understanding language or recognizing images. Machine learning is one way to build such systems. Instead of writing every rule by hand, we show the system examples and adjust it so that its future predictions become more useful.")
    add_body(doc,
        "A model is a mathematical function with adjustable values called parameters. A small classroom model might have a few hundred parameters. A modern language model can have hundreds of millions or billions. It is helpful to imagine the parameters as many connected knobs, but knowledge is not stored as one fact per knob. A behaviour normally depends on patterns spread across many layers and parameters.")

    doc.add_heading("Language models tokens and probabilities", level=2)
    add_body(doc,
        "A language model reads text as tokens. A token may be a whole word, part of a word, punctuation, or a short character sequence. Given the tokens already present, the model assigns a probability to every possible next token. Repeating this next token prediction produces a response.")
    add_table(doc,
        ["Possible next token", "Example teacher probability", "Meaning"],
        [
            ["Paris", "0.90", "The model strongly prefers the correct continuation"],
            ["London", "0.04", "Plausible word form but unlikely factually"],
            ["Berlin", "0.02", "Another low probability alternative"],
            ["All other tokens", "0.04", "The remaining probability mass"],
        ], widths=[2.0, 2.0, 3.0], font_size=9.0)
    add_body(doc,
        "The full list of probabilities is an output distribution. Two models may produce the same final word while assigning very different probabilities. Distribution based comparison can therefore detect drift before an accuracy score changes.")

    doc.add_heading("Pretraining and the base model", level=2)
    add_body(doc,
        "Pretraining creates a broadly capable base model by exposing it to a very large body of text. The resulting model may know language, common facts, patterns of reasoning, and many response styles. An instruction tuned version receives additional training so that it follows user requests more reliably. In this project, the starting checkpoint is our base reference even if its name includes Instruct.")

    doc.add_heading("Training and prompting are different", level=2)
    add_table(doc,
        ["Activity", "Do parameters change", "What happens"],
        [
            ["Ordinary prompting", "No", "The model uses its current parameters to produce an answer"],
            ["Evaluation", "No", "We record answers, probabilities, losses, or other measurements"],
            ["Fine tuning", "Yes", "An optimizer uses a loss signal to update trainable parameters"],
            ["KL anchoring during training", "Yes", "Teacher student disagreement contributes a gradient to the LoRA adapter"],
        ], widths=[1.65, 1.35, 4.0], font_size=8.9)
    add_body(doc,
        "This distinction answers a common question. Giving a trained model one more prompt does not normally change it. The model changes only when we run a training procedure that computes a loss, backpropagates gradients, and applies an optimizer update.")

    doc.add_heading("Fine tuning", level=2)
    add_body(doc,
        "Fine tuning continues training on a narrower dataset. In our project, the target dataset contains mathematical word problems and worked answers. The target loss measures how poorly the student predicts the desired answer tokens. Reducing this loss should improve mathematical performance, but the update can also change responses to unrelated prompts.")

    doc.add_heading("Full fine tuning and LoRA", level=2)
    add_body(doc,
        "Full fine tuning allows all or most original parameters to change. Low Rank Adaptation, known as LoRA, freezes the original weight matrices and adds a smaller trainable update. If the original matrix is W, LoRA uses W plus a product of two small matrices. The rank controls the maximum dimensionality of that update. Far fewer values are trained, which reduces memory and storage requirements [1].")
    add_body(doc,
        "LoRA does not create a separate mathematics brain. The LoRA adapter participates in every forward pass while it is enabled. A gradient produced by one mathematics example changes shared adapter values, and those values may affect science, writing, logic, or other prompts. This is why a frozen backbone can still show capability degradation. Prior work also shows that LoRA often learns less and forgets less than full fine tuning, so comparisons must match target learning rather than count training steps alone [2].")

    part_heading(doc, "Part B Problem and protection")
    doc.add_heading("Capability degradation and catastrophic forgetting", level=2)
    add_body(doc,
        "Suppose a model improves from 20 percent to 45 percent on a mathematics benchmark but falls from 62 percent to 55 percent on a reading task. The seven point reading drop is behavioural evidence of forgetting or capability degradation. It does not by itself prove that information was physically erased from the weights.")
    add_body(doc,
        "Several mechanisms can produce the same visible failure. The model may assign less probability to an old answer, become overly biased toward the new task format, become more sensitive to wording, violate the expected output schema, or lose reliable access to a representation it still contains. Research on implicit inference and spurious forgetting shows that prompt or scoring changes can sometimes recover performance [6, 7].")

    doc.add_heading("The frozen teacher and the trainable student", level=2)
    add_body(doc,
        "We load the same starting checkpoint in two roles. The teacher is frozen and always kept in evaluation mode. It represents the behaviour before mathematics adaptation. The student begins from the same checkpoint, receives a LoRA adapter, and is trained. On target examples, the student follows the mathematics answer. On selected general anchor prompts, the student is encouraged to stay close to the teacher.")
    add_body(doc,
        "The teacher is not assumed to be perfectly correct. It is a behavioural reference. We preserve its distribution on general prompts because our question concerns retention relative to the starting model. We do not use final retention benchmark questions as anchors because that would contaminate the evaluation.")

    doc.add_heading("Anchor prompts", level=2)
    add_body(doc,
        "An anchor prompt is a general input used during training to constrain drift. Examples include a short science explanation, a polite email request, a factual question, or a logical choice. The anchor pool can be unlabeled because the teacher supplies the response distribution. This differs from replay, where a stored prompt and a known target answer are trained again.")
    add_body(doc,
        "We should avoid using mathematics anchors that directly oppose the new task. The purpose is to preserve unrelated general behaviour while allowing the model to improve on the target task. A small number of neutral mathematical or reasoning prompts can be retained only if the research question explicitly includes preservation within the same domain.")

    doc.add_heading("Knowledge distillation and KL divergence", level=2)
    add_body(doc,
        "Knowledge distillation trains a student to imitate signals from a teacher. A hard target says which answer is correct. A soft distribution also shows the teacher's relative preference for alternatives. Kullback Leibler divergence, shortened to KL divergence, measures disagreement between two probability distributions.")
    add_body(doc,
        "Our primary direction is teacher to student KL. In plain text, the calculation is: KL of teacher versus student equals the sum over tokens of teacher probability multiplied by the logarithm of teacher probability divided by student probability. We compute it at each response token position and average across the scored positions.",
        style="Small Note")
    add_table(doc,
        ["Token", "Teacher", "Student after pilot", "What the comparison reveals"],
        [
            ["Paris", "0.90", "0.52", "The student is less confident in the original continuation"],
            ["London", "0.04", "0.15", "An alternative became much more likely"],
            ["Berlin", "0.02", "0.09", "The distribution spread to another alternative"],
            ["Other", "0.04", "0.24", "Much more probability moved outside the original answer"],
        ], widths=[1.35, 1.1, 1.55, 3.0], font_size=8.8)
    add_body(doc,
        "The training objective combines two pressures: total loss equals mathematics target loss plus protection strength multiplied by anchor KL loss. The protection strength, often written as lambda in code, controls the tradeoff. If it is too small, retention may not improve. If it is too large, the student may be prevented from learning mathematics.")

    doc.add_heading("How one anchor update affects other prompts", level=2)
    add_body(doc,
        "A KL update is not stored only for the prompt that produced it. The anchor prompt creates a gradient through the same shared LoRA matrices used for every input. The optimizer changes those matrices, so later predictions on many other prompts may change. We hope the update generalizes to related behaviours, but it can also create side effects. This is the main reason selection quality and held out evaluation matter.")
    add_body(doc,
        "The project does not first discover the exact parameters damaged by mathematics and then prompt those parameters. It discovers vulnerable prompts or behaviours by observing output drift. Parameter localization would require a different mechanistic study using gradients, activations, or causal interventions.")

    part_heading(doc, "Part C Proposed method")
    add_figure(doc, WORKFLOW_PNG, 6.25,
               "Figure 1  Pilot guided selection is separate from every final training condition",
               "A seven stage vertical workflow showing the frozen reference, pilot LoRA run, damage scoring, anchor selection, pilot reset, main training, and matched evaluation.")

    doc.add_heading("Step one establish the reference", level=2)
    add_body(doc,
        "Load the untouched instruction model. Evaluate it on the mathematics target validation set and on candidate retention tasks. Record item level predictions, output format failures, and likelihood scores. Retention tasks where the base model is at or near chance should not be primary benchmarks because there is little correct behaviour available to forget.")

    doc.add_heading("Step two create a disjoint anchor pool", level=2)
    add_body(doc,
        "Collect about 2,000 general instruction prompts from a public source such as Dolly 15k or another permitted prompt dataset. Remove empty, unsafe, extremely long, mathematics target like, and duplicate items. Use exact matching and embedding similarity to remove overlap with target validation and final evaluation questions. Store the cleaned pool before looking at final results.")

    doc.add_heading("Step three run a short pilot LoRA", level=2)
    add_body(doc,
        "Create a temporary student from the base checkpoint. Train a LoRA adapter for roughly 5 to 10 percent of the planned target training steps on a fixed mathematics subset. The pilot should be long enough to create detectable output drift but short enough to remain inexpensive. Save the precise pilot data indices, seed, hyperparameters, and checkpoint.")

    doc.add_heading("Step four create fixed teacher continuations", level=2)
    add_body(doc,
        "For every unlabeled candidate prompt, let the frozen teacher generate one short deterministic continuation and save it. Then force both teacher and pilot to score exactly that same prompt continuation sequence. This avoids comparing two generated strings with different lengths or token alignments. KL is calculated only on the saved response token positions, not on padding.")

    doc.add_heading("Step five calculate the damage score", level=2)
    add_body(doc,
        "For each candidate, average teacher to pilot KL across the scored response positions. A high value means the pilot changed its distribution strongly on that prompt. This is a behavioural vulnerability score, not proof that the prompt would become incorrect in the final model.")
    add_table(doc,
        ["Candidate prompt", "Illustrative score", "Interpretation"],
        [
            ["Write a friendly greeting", "0.02", "The pilot remains close to the teacher"],
            ["Explain photosynthesis simply", "0.31", "Moderate early drift"],
            ["Choose the logically valid conclusion", "0.74", "Large early drift and a strong protection candidate"],
        ], widths=[3.25, 1.35, 2.4], font_size=8.9)
    add_body(doc,
        "The primary score should be chosen before the main runs. Alternative scores such as teacher negative log likelihood increase, answer margin change, or entropy adjusted drift belong in a later ablation only if the central matrix is complete.")

    doc.add_heading("Step six select four equal size anchor sets", level=2)
    add_body(doc,
        "Semantic means meaning. Two sentences are semantically similar when they express similar ideas even if they use different words. For example, Write a polite apology and Draft a courteous message saying sorry are semantically close. An embedding model converts each prompt into a numerical vector. Similar meanings tend to have nearby vectors, which can be grouped with a clustering algorithm.")
    add_table(doc,
        ["Selector", "Rule", "What it tests", "Main weakness"],
        [
            ["Random", "Sample anchors uniformly from the filtered pool", "Whether ordinary anchoring is already sufficient", "Can waste budget on duplicates or unaffected prompts"],
            ["Diversity", "Cluster prompt embeddings and sample across clusters", "Whether broad meaning coverage matters", "May protect stable prompts"],
            ["Damage", "Choose prompts with the highest pilot drift", "Whether vulnerability predicts useful protection", "Top prompts may be redundant or noisy"],
            ["Diverse damage", "Retain high drift candidates, cluster them, then choose across clusters", "Whether vulnerability and coverage work together", "Has a pilot and clustering cost"],
        ], widths=[1.25, 2.2, 2.0, 1.55], font_size=8.3)
    add_figure(doc, SELECTION_PNG, 6.35,
               "Figure 2  Conceptual comparison of four selectors with the same anchor count",
               "Four scatter plots show random, diversity, damage, and diverse damage selection. Points close together represent similar meanings, larger points have higher damage scores, and selected anchors are circled in red.")

    doc.add_heading("Step seven discard the pilot and reset", level=2)
    add_body(doc,
        "The pilot is diagnostic only. Delete it from the final comparison and initialize every main student from the exact same base checkpoint. The pilot must not continue into Damage KL or Diverse Damage KL, because those conditions would otherwise receive extra mathematics training.")

    doc.add_heading("Step eight train the main conditions", level=2)
    add_body(doc,
        "Each main run receives identical target data order, target token budget, LoRA configuration, optimizer settings, and checkpoint schedule for its seed. On anchor steps, the selected general prompts contribute KL loss. Random, diversity, damage, and diverse damage conditions receive the same number of scored response tokens and the same number of teacher passes.")
    add_body(doc,
        "Target and anchor batches can be alternated, such as four target steps followed by one anchor step, or combined in the same optimizer step. Alternation is easier to inspect. Whichever schedule we choose must remain identical across all KL conditions. The task only condition receives target training without anchor steps; its number of target updates stays equal to the others.")

    doc.add_heading("Step nine evaluate and interpret", level=2)
    add_body(doc,
        "Evaluate saved checkpoints on target validation and select comparable checkpoints without using the test set. Run final target and retention evaluation once the selection rule is frozen. Then apply the recovery battery to items that became incorrect. Report learning, retention, recoverability, wall time, and selection overhead.")

    doc.add_heading("Conceptual pseudocode", level=2)
    code = (
        "teacher = frozen_base_model()\n"
        "pilot = add_lora(fresh_base_model())\n"
        "train_briefly(pilot, math_subset)\n\n"
        "for prompt in candidate_pool:\n"
        "    continuation = teacher.generate(prompt)\n"
        "    score[prompt] = token_average_KL(teacher, pilot, prompt, continuation)\n\n"
        "anchor_sets = select_random_diverse_damage_and_diverse_damage(score)\n"
        "discard(pilot)\n\n"
        "for condition in main_conditions:\n"
        "    student = add_lora(fresh_base_model())\n"
        "    repeat training schedule:\n"
        "        task_loss = supervised_math_loss(student, target_batch)\n"
        "        anchor_loss = teacher_student_KL(student, condition_anchor_batch)\n"
        "        update_lora(task_loss + protection_strength * anchor_loss)"
    )
    add_body(doc, code, style="Code Block", keep=True)

    part_heading(doc, "Part D Experimental design")
    doc.add_heading("Central experimental conditions", level=2)
    add_table(doc,
        ["Condition", "Training data", "Retention objective", "Scientific role"],
        [
            ["Base", "None", "None", "Untouched evaluation reference only"],
            ["Task only LoRA", "GSM8K", "None", "Measures ordinary learning and forgetting"],
            ["Random KL", "GSM8K plus random anchors", "KL to the base teacher", "Plain anchoring baseline"],
            ["Diversity KL", "GSM8K plus broad anchors", "KL to the base teacher", "Meaning coverage baseline"],
            ["Damage KL", "GSM8K plus highest drift anchors", "KL to the base teacher", "Tests the pilot damage signal"],
            ["Diverse Damage KL", "GSM8K plus damaged broad anchors", "KL to the base teacher", "Primary proposed method"],
        ], widths=[1.3, 2.0, 1.75, 1.95], font_size=8.4)
    add_body(doc,
        "The base model is evaluated but not trained. The five training conditions repeated with two seeds produce ten central runs. The shared pilot is an additional diagnostic run. Full fine tuning, second models, labeled replay, large rank sweeps, and multiple target tasks are outside the core plan so that the team can finish a clean study.")

    doc.add_heading("Data plan", level=2)
    add_table(doc,
        ["Data component", "Recommended source and size", "Use", "Separation rule"],
        [
            ["Target training", "2,000 to 4,000 GSM8K training examples", "Mathematics LoRA adaptation", "No test items and fixed training indices"],
            ["Target validation", "A held out subset from training data", "Checkpoint selection and target matching", "Never used for gradient updates"],
            ["Target test", "Official GSM8K test split", "Final target result", "Evaluate after decisions are frozen"],
            ["Anchor candidates", "About 2,000 filtered general prompts", "Selection and KL training", "No final benchmark questions or near duplicates"],
            ["Retention development", "Small development subsets", "Prompt templates and parser debugging", "Different items from final retention tests"],
            ["Retention test", "Qualified subsets of general benchmarks", "Final forgetting and recovery measures", "Never used for anchors or hyperparameter tuning"],
        ], widths=[1.45, 2.15, 1.65, 1.75], font_size=8.0)

    doc.add_heading("Benchmark qualification", level=2)
    add_body(doc,
        "Before training, measure the base model on candidate retention tasks. Keep tasks that have enough correct base items to support item transition analysis. A task at chance can still provide likelihood information, but its accuracy drop cannot show much forgetting. The final suite should cover different behaviours rather than many versions of the same multiple choice task.")
    add_table(doc,
        ["Capability", "Candidate measure", "Why include it"],
        [
            ["Reading and factual judgment", "BoolQ or selected nonmath MMLU subjects", "Tests general question answering"],
            ["Commonsense completion", "HellaSwag", "Tests contextual continuation"],
            ["Science reasoning", "ARC Challenge or OpenBookQA", "Tests reasoning outside the target domain"],
            ["Instruction compliance", "Compact IFEval subset and format checks", "Detects style and schema degradation"],
            ["General language modeling", "Held out text negative log likelihood", "Provides a sensitive distributional measure"],
        ], widths=[1.85, 2.55, 2.6], font_size=8.7)

    doc.add_heading("Fixed anchor token and training budget", level=2)
    add_body(doc,
        "An anchor count alone is not enough because prompts and responses have different lengths. We should cap or batch by scored response tokens. As an illustrative design, 128 anchors with up to 96 scored response tokens repeated four times would provide at most 49,152 anchor tokens. The final number should be chosen after a small timing test and then frozen.")
    add_bullets(doc, [
        "Give every KL selector the same anchor token allowance, anchor step schedule, and teacher scoring method.",
        "Give every trained condition the same target examples, target tokens, and optimizer update opportunity.",
        "Record main training GPU hours for every condition.",
        "Report the shared pilot and selection cost separately, because damage based selection has overhead that random selection does not.",
    ])
    add_body(doc,
        "We should not claim equal total compute if the pilot overhead is excluded. The scientific comparison can hold main training constant, while a separate cost effectiveness table reports end to end GPU hours including selection.")

    doc.add_heading("Target performance matching", level=2)
    add_body(doc,
        "A protection method can appear to forget less because it also blocks mathematics learning. We therefore save checkpoints throughout training and compare checkpoints whose GSM8K validation scores are within a predeclared tolerance, such as one or two percentage points. If the learning ranges do not overlap, we show the full mathematics gain versus forgetting curve instead of forcing an invalid match.")
    add_numbers(doc, [
        "Choose the target validation metric and matching tolerance before opening final test results.",
        "For each condition and seed, identify checkpoints within the shared target performance range.",
        "Compare retention at those matched checkpoints and display the full tradeoff frontier.",
        "Use the final target test split only for the locked checkpoints.",
    ])

    doc.add_heading("Recovery aware measurement", level=2)
    add_body(doc,
        "A standard prompt may underestimate what a model can still express. For each retention item that was correct before training and wrong afterward, we apply a fixed recovery battery. The battery is defined on development items and then frozen.")
    add_table(doc,
        ["Recovery stage", "Procedure", "Interpretation"],
        [
            ["Standard", "Use the original benchmark prompt and normal answer parser", "Raw observed performance"],
            ["Template", "Use five meaning preserving prompt wordings", "Recovery suggests prompt sensitivity"],
            ["Likelihood", "For multiple choice tasks score option likelihoods directly", "Recovery suggests generation or format failure"],
            ["Few shot", "Add fixed demonstrations from a development split", "Recovery suggests task inference or access failure"],
            ["Residual", "Count failures that remain after the registered probes", "Stronger behavioural loss but not proof of physical erasure"],
        ], widths=[1.15, 3.35, 2.5], font_size=8.4)
    add_body(doc,
        "Recovery evaluation does not update model parameters. It changes only the test prompt or scoring method. We will report prompt recoverable, likelihood recoverable, few shot recoverable, and residual losses. We will use the phrase residual behavioural loss rather than proven knowledge erasure.")

    doc.add_heading("Metrics", level=2)
    add_table(doc,
        ["Question", "Metric", "How it is calculated"],
        [
            ["Did mathematics improve", "GSM8K numeric exact match", "Extract and compare the final numerical answer"],
            ["Did target modeling improve", "Target negative log likelihood", "Average answer token loss on held out target data"],
            ["How much general performance fell", "Absolute retention drop", "Base accuracy minus adapted accuracy for each task"],
            ["Which items changed", "Item transitions", "Retained, forgotten, newly correct, or still incorrect"],
            ["How sensitive is the result", "Template mean and worst case", "Aggregate across five fixed prompt variants"],
            ["How much loss remains", "Residual drop", "Failure remaining after all registered recovery probes"],
            ["Was the anchor budget efficient", "Retention gain per 100,000 anchor tokens", "Improvement over task only divided by protection tokens"],
            ["Was the method practical", "GPU hours and peak memory", "Main training and selection overhead reported separately"],
        ], widths=[1.85, 2.0, 3.15], font_size=8.1)

    doc.add_heading("Item transition accounting", level=2)
    add_table(doc,
        ["Before fine tuning", "After fine tuning", "Category", "Meaning"],
        [
            ["Correct", "Correct", "Retained", "The original behaviour remained successful"],
            ["Correct", "Incorrect", "Forgotten", "The adaptation introduced a failure"],
            ["Incorrect", "Correct", "Newly correct", "Positive backward transfer occurred"],
            ["Incorrect", "Incorrect", "Still incorrect", "The item cannot demonstrate forgetting"],
        ], widths=[1.35, 1.35, 1.45, 2.85], font_size=8.8)

    doc.add_heading("Hypotheses", level=2)
    add_table(doc,
        ["ID", "Falsifiable hypothesis"],
        [
            ["H1", "Damage KL will reduce general capability loss more than Random KL at the same anchor token budget and matched target score"],
            ["H2", "Diverse Damage KL will outperform simple Damage KL when the highest drift prompts contain semantic duplicates"],
            ["H3", "The winning selector will improve likelihood and multi template retention, not only generated exact match"],
            ["H4", "Some raw forgetting will be recoverable through prompt or scoring changes"],
            ["H5", "Excessive protection strength will improve retention while lowering target learning"],
        ], widths=[0.65, 6.35], font_size=8.9)

    doc.add_heading("Primary analysis and statistics", level=2)
    add_body(doc,
        "The primary comparison is Diverse Damage KL versus Random KL at the chosen middle anchor budget and matched GSM8K validation performance. We will show each seed separately and aggregate item level differences with paired bootstrap confidence intervals. Two seeds are a practical course minimum, so we will avoid strong claims about training randomness and focus on effect size, consistency, and item level uncertainty.")
    add_bullets(doc, [
        "Report per task scores as well as a clearly defined average. Do not hide opposite task effects inside one mean.",
        "Use paired resampling because every model is evaluated on the same items.",
        "Report confidence intervals and the number of base correct items for each retention task.",
        "Treat secondary pairwise comparisons as descriptive unless a correction plan is registered.",
        "Inspect failures qualitatively after the quantitative analysis rule is fixed.",
    ])

    doc.add_heading("Expected figures and tables", level=2)
    add_table(doc,
        ["Output", "What the reader should learn"],
        [
            ["Target gain versus raw forgetting plot", "Whether a method protects capability without blocking mathematics learning"],
            ["Target gain versus residual forgetting plot", "Whether protection remains after recovery controls"],
            ["Item transition chart", "How many base correct items were retained, forgotten, or recovered"],
            ["Anchor efficiency chart", "Retention benefit per fixed anchor token budget"],
            ["Selector coverage plot", "Whether damage only selection concentrates in a few meaning clusters"],
            ["Cost table", "Main training time, pilot overhead, memory, and cached storage"],
        ], widths=[2.65, 4.35], font_size=8.8)

    part_heading(doc, "Part E Implementation")
    doc.add_heading("Model state during training", level=2)
    add_table(doc,
        ["Component", "State", "Purpose"],
        [
            ["Base backbone in the student", "Frozen", "Provides the shared pretrained representation"],
            ["LoRA matrices in the student", "Trainable", "Learn the target task and respond to KL protection"],
            ["Teacher behaviour", "Frozen with dropout disabled", "Provides the original output distribution"],
            ["Pilot adapter", "Temporary", "Creates the vulnerability signal and is then discarded"],
            ["Optimizer", "Updates student LoRA only", "Applies gradients from task and anchor losses"],
        ], widths=[2.05, 2.0, 2.95], font_size=8.8)

    doc.add_heading("Efficient teacher computation", level=2)
    add_body(doc,
        "A separate full teacher copy is conceptually simple but may waste GPU memory. Because the student backbone is the frozen base, one implementation can disable the LoRA adapter for the teacher forward pass and enable it for the student pass. The teacher pass must not build gradients. The code should verify that teacher outputs remain identical across training steps.")
    add_body(doc,
        "For repeatable anchor scoring, cache each teacher continuation. We can then compute teacher logits online or cache a compressed top token distribution. Full vocabulary logits for every token may be too large. A top k approximation should preserve the teacher probability mass outside the saved tokens in an other bucket and must be used consistently for every condition.")

    doc.add_heading("Numerical and implementation checks", level=2)
    add_bullets(doc, [
        "Apply log softmax in stable precision before KL calculation and mask padding positions.",
        "Confirm that teacher probabilities do not receive gradients and teacher dropout is disabled.",
        "Unit test that KL is close to zero when teacher and student logits are identical.",
        "Check that the loss grows when a controlled perturbation changes the student logits.",
        "Log target loss and anchor loss separately so that one cannot silently dominate.",
        "Verify that every KL condition consumes the same counted anchor tokens.",
        "Hash all split files and record dataset versions before the first final run.",
    ])

    doc.add_heading("Recommended repository structure", level=2)
    add_table(doc,
        ["Path", "Purpose"],
        [
            ["configs/", "One frozen configuration per condition and seed"],
            ["data_manifests/", "Dataset IDs and hashes rather than uncontrolled raw copies"],
            ["src/train_lora.py", "Target only and KL training loop"],
            ["src/score_damage.py", "Teacher versus pilot prompt scoring"],
            ["src/select_anchors.py", "Random, diversity, damage, and diverse damage selection"],
            ["src/evaluate.py", "Target, retention, item transition, and recovery evaluation"],
            ["tests/", "Loss, masking, budget, and split checks"],
            ["results/", "Item level outputs and run summaries"],
            ["reports/", "Plots, tables, paper, and presentation"],
        ], widths=[2.05, 4.95], font_size=8.6)

    doc.add_heading("Run naming and records", level=2)
    add_body(doc,
        "Use a predictable identifier such as qwen05b_damagekl_seed17. Every run record should contain the base checkpoint revision, LoRA rank and alpha, targeted modules, learning rate, optimizer, batch and accumulation settings, maximum sequence length, target and anchor tokens, protection strength, seed, hardware, wall time, peak memory, and checkpoint evaluation scores.")
    add_table(doc,
        ["Artifact", "Example filename", "Why it exists"],
        [
            ["Base evaluation", "base_eval_items.parquet", "Starting answers and probabilities for item transitions"],
            ["Pilot scores", "pilot_anchor_scores.parquet", "Prompt IDs, damage values, lengths, and clusters"],
            ["Anchor manifest", "anchors_diverse_damage_seed17.jsonl", "Exact fixed prompts used in one condition"],
            ["Run configuration", "qwen05b_damagekl_seed17.yaml", "Complete training settings"],
            ["Evaluation outputs", "qwen05b_damagekl_seed17_items.parquet", "Per item answers, scores, and recovery category"],
            ["Summary", "central_results.csv", "One row per condition, seed, checkpoint, and metric"],
        ], widths=[1.45, 2.5, 3.05], font_size=8.4)

    doc.add_heading("Contamination prevention", level=2)
    add_body(doc,
        "The anchor pool, recovery examples, target validation data, and final evaluation data must remain disjoint. Exact string checks catch direct duplicates. Normalized checks remove punctuation and case differences. Embedding similarity helps flag near duplicates for manual review. We should publish the IDs and hashes of every split so that the separation can be audited.")

    doc.add_heading("Reproducibility checklist", level=2)
    add_bullets(doc, [
        "Pin model, tokenizer, dataset, Transformers, PEFT, and evaluation library versions.",
        "Save configurations rather than relying on notebook state.",
        "Keep target and anchor data order reproducible for each seed.",
        "Save regular checkpoints so that target matching remains possible.",
        "Retain item level predictions rather than only final averages.",
        "Run one clean reproduction from the documented command on a second team member's machine.",
        "Record failed runs and changes to the preregistered plan in a deviation log.",
    ])

    part_heading(doc, "Part F Execution")
    doc.add_heading("Eight week plan", level=2)
    add_table(doc,
        ["Week", "Main work", "Exit criterion"],
        [
            ["1", "Select the model, qualify retention benchmarks, create disjoint manifests, and reproduce base evaluation", "The same base results reproduce on two machines"],
            ["2", "Implement task only LoRA and save frequent checkpoints", "Mathematics improves and at least one sensitive retention signal changes"],
            ["3", "Implement teacher scoring and Random KL with loss and budget tests", "KL is numerically stable and the random baseline completes"],
            ["4", "Run the shared pilot, score candidates, cluster embeddings, and freeze four anchor manifests", "Selectors pass size, length, overlap, and diversity checks"],
            ["5", "Run all five central conditions for the first seed", "Checkpoints, logs, and item outputs are complete"],
            ["6", "Repeat the central matrix for the second seed and target match checkpoints", "The primary comparison has two reproducible seeds"],
            ["7", "Run recovery evaluation, paired bootstrap analysis, cost accounting, and qualitative review", "All planned tables and figures can be generated from saved outputs"],
            ["8", "Write the report, build the presentation, package code, and reproduce one run end to end", "Another teammate follows the documentation without hidden steps"],
        ], widths=[0.55, 3.7, 2.75], font_size=8.0)

    doc.add_heading("Team responsibilities", level=2)
    add_table(doc,
        ["Owner", "Primary responsibility", "Required handoff"],
        [
            ["Member A", "Data manifests, contamination checks, benchmark qualification, and base evaluation", "Frozen IDs, hashes, prompts, parsers, and baseline results"],
            ["Member B", "LoRA training loop, KL loss, teacher computation, checkpoints, and run logging", "Tested trainer plus complete configurations and logs"],
            ["Member C", "Pilot run, damage scoring, embeddings, clustering, and anchor manifests", "Reproducible selector outputs with diagnostics"],
            ["Member D", "Recovery battery, statistics, plots, cost accounting, and report integration", "Item level analyses, confidence intervals, figures, and limitations"],
        ], widths=[1.0, 3.5, 2.5], font_size=8.4)
    add_body(doc,
        "For a three person team, Member A can also own the recovery datasets, Member B can own cost logging, and Member C can own statistics and plotting. Each person should review one component owned by someone else. No result should depend on an undocumented local notebook.")

    weekly_heading = doc.add_heading("Weekly team routine", level=2)
    weekly_heading.paragraph_format.page_break_before = True
    add_bullets(doc, [
        "Freeze one written experiment plan at the start of each week.",
        "Use a shared run table with status, owner, configuration hash, output path, and notes.",
        "Hold a short midweek check for failures that affect other members.",
        "End each week with one reproducible artifact and a decision about the next dependency.",
        "Do not launch a new ablation while a core run or evaluation bug remains unresolved.",
    ])

    doc.add_heading("Pilot decision rules", level=2)
    add_body(doc,
        "Week two is the first major gate. The team should predefine a reasonable minimum target improvement and a detectable retention change. An illustrative gate is at least a five point GSM8K improvement and either a two point drop on one qualified benchmark or a clear held out likelihood shift. These are planning values, not universal scientific thresholds. If the base starts near the floor, the team should change the model or benchmark before investing in the full matrix.")

    doc.add_heading("Risks and mitigations", level=2)
    add_table(doc,
        ["Risk", "Why it matters", "Response"],
        [
            ["No measurable forgetting", "Selectors cannot improve a signal that is absent", "Increase target steps gradually, use a more sensitive likelihood metric, or change the qualified retention suite"],
            ["No mathematics learning", "Low forgetting may only reflect failed adaptation", "Check formatting, learning rate, data, and choose a model with nonfloor target capacity"],
            ["Pilot scores are unstable", "Selected anchors may reflect noise", "Measure score agreement across pilot checkpoints or one additional pilot seed before main runs"],
            ["High damage prompts are duplicates", "Damage selection wastes budget", "Use diverse damage selection and report cluster concentration"],
            ["KL blocks target learning", "The method appears safe for the wrong reason", "Lower protection strength and rely on matched checkpoints and tradeoff curves"],
            ["Teacher cache is too large", "Storage or input output becomes a bottleneck", "Cache continuations and top token probabilities or compute teacher logits online"],
            ["Evaluation leakage", "Retention results become invalid", "Freeze disjoint manifests and audit exact and near duplicates"],
            ["Two seeds disagree", "The conclusion may depend on training randomness", "Report both, avoid overclaiming, and spend any remaining compute on the primary pair"],
        ], widths=[1.55, 2.35, 3.1], font_size=7.9)

    doc.add_heading("Minimum viable study", level=2)
    add_body(doc,
        "If time or compute becomes tight, retain the scientific heart of the project. Use one 0.5B model, GSM8K, Task only LoRA, Random KL, Damage KL, and Diverse Damage KL, two seeds, three qualified retention tasks, forced choice scoring, and five prompt templates. Diversity KL is the first condition to remove because Diverse Damage KL already uses semantic clustering.")

    doc.add_heading("Fallbacks that preserve the research question", level=2)
    add_table(doc,
        ["Problem", "Fallback", "Limitation to report"],
        [
            ["Exact KL is too expensive", "Match only teacher top token probabilities plus an other mass bucket", "This approximates full distribution KL"],
            ["Teacher forward passes are too slow", "Cache compressed teacher distributions once", "Storage rises and tokenization must remain fixed"],
            ["Forgetting stays near zero", "Use a longer target adaptation or a more format distant narrow task", "The changed protocol must be recorded before final testing"],
            ["GSM8K is too hard for the model", "Use an easier arithmetic subset or a 1.5B model if hardware permits", "The new target changes external validity and compute"],
            ["Clustering quality is poor", "Use a simple farthest first diversity rule on normalized embeddings", "Coverage is geometric rather than cluster based"],
        ], widths=[1.7, 3.0, 2.3], font_size=8.2)

    doc.add_heading("Work deliberately excluded from the core", level=2)
    add_body(doc,
        "The core project will not attempt a broad comparison of O LoRA, CURLoRA, OPLoRA, full fine tuning, multiple ranks, multiple base models, and several target tasks. Those are reasonable future extensions, but including them now would weaken the controls and analysis. The course project should complete one contribution and its necessary baselines.")

    part_heading(doc, "Part G Interpretation and communication")
    doc.add_heading("How to interpret every likely outcome", level=2)
    add_table(doc,
        ["Observed result", "Supported interpretation", "What we should not claim"],
        [
            ["Diverse Damage beats Random at matched target learning", "Pilot vulnerability plus coverage selected a more useful small anchor set", "That the method works for every model, task, or budget"],
            ["Damage beats Random but Diverse Damage does not", "High drift mattered, while clustering removed useful concentration", "That diversity never helps"],
            ["Diversity beats Damage", "Broad coverage mattered more than the pilot score in this setting", "That the pilot idea is universally wrong"],
            ["All KL methods tie", "The pool may be redundant, the budget may be large enough, or the pilot score may be uninformative", "That selection can never matter"],
            ["KL retains more but learns less mathematics", "Protection created a stability plasticity tradeoff", "That KL preserved more at equal learning"],
            ["Raw gains vanish after recovery controls", "The method mainly preserved prompt alignment or output format", "That it preserved deeper capability"],
            ["Residual gains remain", "The method preserved behaviour across stronger elicitation tests", "That neural knowledge storage was proven unchanged"],
        ], widths=[2.15, 3.15, 1.7], font_size=7.9)

    doc.add_heading("What a strong negative result looks like", level=2)
    add_body(doc,
        "A null result can still answer a useful question if the comparison is fair. If random and damage aware anchors tie within narrow confidence intervals at matched target learning, we can conclude that the added pilot and selection cost did not provide a detectable benefit in the tested small model setting. The paper should then explain score stability, coverage, budget size, and task sensitivity rather than searching for an unplanned positive subgroup.")

    doc.add_heading("Questions teammates and the professor may ask", level=2)
    faq_rows = [
        ["Are we finding which model parameters mathematics damaged", "No. We measure which prompts show early output distribution drift. Parameter localization is a separate mechanistic problem."],
        ["Does every prompt permanently change the model", "No. Ordinary prompting and evaluation do not update parameters. Only training steps with gradients and an optimizer change the LoRA adapter."],
        ["Is KL computed for every user prompt", "No. During training it is computed only on scheduled anchor prompts. During ordinary use after training there is no teacher comparison."],
        ["Does an anchor update affect only that anchor", "No. It updates shared LoRA matrices and can change responses to related and unrelated inputs."],
        ["Why discard the pilot", "The pilot is a measuring instrument. Keeping it would give damage based conditions extra target training and invalidate the comparison."],
        ["Why use unlabeled anchors", "The frozen teacher supplies a full soft distribution, so labels are unnecessary and the pool can be broader."],
        ["Why not choose only the highest damage scores", "Those prompts may be near duplicates. Diverse Damage tries to preserve high vulnerability while covering different meanings."],
        ["Why match mathematics performance", "A method that learns less often forgets less. Matching isolates retention from reduced adaptation."],
        ["Is KL anchoring itself novel", "No. The proposed contribution is pilot guided anchor selection plus a fixed budget and recovery aware controlled evaluation."],
        ["What if the method fails", "A careful null or negative result identifies when added selector complexity is not justified and remains valid course research."],
    ]
    add_table(doc, ["Question", "Answer"], faq_rows, widths=[2.55, 4.45], font_size=8.4)

    doc.add_heading("Glossary", level=2)
    glossary = [
        ["Anchor prompt", "A general prompt used during training to limit drift from the original model"],
        ["Base model", "The untouched starting checkpoint used as the reference"],
        ["Candidate pool", "The filtered collection from which anchors are selected"],
        ["Catastrophic forgetting", "Loss of earlier task performance after learning something new"],
        ["Damage score", "A pilot based measure of how much output behaviour changed on one candidate prompt"],
        ["Embedding", "A numerical vector that represents aspects of a prompt's meaning"],
        ["Fine tuning", "Additional training that adapts a pretrained model to a narrower objective"],
        ["Gradient", "A signal indicating how trainable values should change to reduce loss"],
        ["KL divergence", "A measure of disagreement between two probability distributions"],
        ["LoRA", "A method that freezes the original model and trains small low rank update matrices"],
        ["Parameter", "An adjustable numerical value inside a model"],
        ["Pilot run", "A short temporary adaptation used to estimate vulnerability"],
        ["Recovery battery", "Several fixed prompting and scoring tests that probe whether an apparent loss is recoverable"],
        ["Residual loss", "Failure that remains after the registered recovery probes"],
        ["Seed", "A number controlling randomized parts of a run so they can be repeated"],
        ["Semantic similarity", "Similarity in meaning rather than exact wording"],
        ["Student", "The model with the trainable LoRA adapter"],
        ["Target matching", "Comparing checkpoints with approximately equal target task performance"],
        ["Teacher", "The frozen starting model that supplies reference distributions"],
        ["Token", "A word, word piece, symbol, or punctuation unit processed by the model"],
    ]
    add_table(doc, ["Term", "Meaning in this project"], glossary, widths=[1.65, 5.35], font_size=8.3)

    doc.add_heading("One minute explanation for teammates", level=2)
    add_body(doc,
        "We start with a small language model that can do many things. We specialize it in mathematics using LoRA, but that specialization may change unrelated abilities. We keep the original model frozen as a teacher. A short temporary mathematics run first shows us which general prompts change most. We then throw that temporary model away, restart from the original model, and train separate students. Each student gets the same mathematics training, but the protective prompts are chosen randomly, for broad meaning coverage, by high pilot damage, or by both damage and diversity. On those prompts, KL divergence encourages the student to behave like the frozen teacher. We compare the methods at similar mathematics performance and use multiple prompt formats to see whether failures are recoverable. The research question is whether a limited protection budget works better when it is spent on vulnerable behaviours rather than random prompts.",
        style="Lead")

    doc.add_heading("Professor ready abstract", level=2)
    add_body(doc,
        "Fine tuning a language model on a narrow task can improve target performance while degrading unrelated capabilities. Parameter efficient adaptation reduces the number of trainable parameters but does not eliminate this problem. We propose a controlled study of damage aware KL anchoring for a small instruction model adapted to mathematical word problems with LoRA. A frozen base model supplies reference output distributions on an unlabeled pool of general prompts. A short pilot adaptation estimates prompt vulnerability from base to pilot token distribution drift. The pilot is then discarded, and fresh students are trained with equal target data and equal anchor token budgets using random, diversity based, damage based, or diversity filtered damage based anchors. We compare these conditions with task only LoRA at target matched checkpoints. Evaluation includes target exact match, general capability retention, item level transitions, forced choice likelihood, prompt templates, residual behavioural loss, and compute cost. The primary hypothesis is that diverse damage aware anchors preserve more held out capability than random anchors without reducing target learning. The contribution is the controlled use of early behavioural damage for unlabeled anchor selection and recovery aware measurement, rather than KL regularization itself.")

    doc.add_heading("Proposal checklist", level=2)
    add_bullets(doc, [
        "State one primary research question and the Diverse Damage versus Random primary comparison.",
        "Describe KL, LoRA, targeted selection, and recovery methods as prior ideas rather than claiming them as inventions.",
        "Define all data splits and prevent anchor evaluation overlap.",
        "Specify the anchor token budget, target token budget, checkpoint schedule, and compute accounting.",
        "Predeclare the target matching tolerance and recovery battery.",
        "Name the ten central runs and the shared pilot.",
        "Explain how negative results will be interpreted.",
        "Repeat a focused literature search immediately before submission.",
    ])

    references_heading = doc.add_heading("References", level=1)
    references_heading.paragraph_format.page_break_before = True
    add_body(doc,
        "These sources establish the main ingredients and the closest research context. The project proposal should cite the papers it actually discusses and verify bibliographic details before submission.", style="Small Note")
    refs = [
        (1, "LoRA: Low-Rank Adaptation of Large Language Models", "ICLR 2022", "https://openreview.net/forum?id=nZeVKeeFYf9", "Introduces the low rank adaptation method used for the student model"),
        (2, "LoRA Learns Less and Forgets Less", "TMLR 2024", "https://openreview.net/forum?id=aloEru2qCG", "Frames LoRA retention as a learning and forgetting tradeoff"),
        (3, "An Efficient Rehearsal Scheme for Catastrophic Forgetting Mitigation during Multi-stage Fine-tuning", "Findings of NAACL 2025", "https://aclanthology.org/2025.findings-naacl.138/", "Provides a close precedent for damage aware selection of rehearsal examples"),
        (4, "Demystifying Language Model Forgetting with Low-rank Example Associations", "NeurIPS 2025", "https://proceedings.neurips.cc/paper_files/paper/2025/hash/06872e1e6d11baf2ae27285c50132f4f-Abstract-Conference.html", "Predicts example level forgetting and motivates careful novelty claims"),
        (5, "Context-Free Synthetic Data Mitigates Forgetting", "arXiv 2025 and revised 2026", "https://arxiv.org/abs/2505.13811", "Studies context free generations for approximating KL and mitigating forgetting"),
        (6, "Spurious Forgetting in Continual Learning of Language Models", "ICLR 2025", "https://proceedings.iclr.cc/paper_files/paper/2025/hash/a774503daed55eb53c634847ae071ec7-Abstract-Conference.html", "Shows that measured forgetting can include an alignment loss component"),
        (7, "Understanding Catastrophic Forgetting in Language Models via Implicit Inference", "ICLR 2024", "https://proceedings.iclr.cc/paper_files/paper/2024/hash/692ae28fda9bfbde7c01b13bf5a03395-Abstract-Conference.html", "Demonstrates that prompting can recover some apparently lost behaviour"),
        (8, "Mapping Post-Training Forgetting in Language Models at Scale", "ICLR 2026", "https://proceedings.iclr.cc/paper_files/paper/2026/hash/bc1c5e5fb8ed1ef9b9b5abced2022e40-Abstract-Conference.html", "Supports item level transition analysis"),
        (9, "Training Verifiers to Solve Math Word Problems", "arXiv 2021", "https://arxiv.org/abs/2110.14168", "Introduces the GSM8K dataset"),
        (10, "Measuring Massive Multitask Language Understanding", "ICLR 2021", "https://openreview.net/forum?id=d7KBjmI3GmQ", "Introduces MMLU as a broad knowledge benchmark"),
        (11, "HellaSwag: Can a Machine Really Finish Your Sentence?", "ACL 2019", "https://aclanthology.org/P19-1472/", "Introduces a commonsense completion benchmark"),
        (12, "Instruction-Following Evaluation for Large Language Models", "arXiv 2023", "https://arxiv.org/abs/2311.07911", "Introduces IFEval and its verifiable instruction constraints"),
        (13, "Mitigating Catastrophic Forgetting in Large Language Models with Self-Synthesized Rehearsal", "ACL 2024", "https://aclanthology.org/2024.acl-long.77/", "Establishes a strong synthetic rehearsal direction"),
        (14, "LoRA vs Full Fine-tuning: An Illusion of Equivalence", "NeurIPS 2025", "https://papers.nips.cc/paper_files/paper/2025/hash/ff541950d1e885af90f523571564a401-Abstract-Conference.html", "Shows that spectral behaviour and hyperparameters affect LoRA comparisons"),
    ]
    for ref in refs:
        add_reference(doc, *ref)

    doc.add_heading("Final project statement", level=1)
    add_body(doc,
        "We will investigate whether early behavioural drift can guide a limited KL protection budget during LoRA mathematics adaptation. A short pilot identifies vulnerable general prompts, but it is discarded before the final runs. Fresh students then compare random, diversity based, damage based, and diverse damage based anchors under controlled budgets. We will judge success at matched mathematics performance and distinguish prompt recoverable failures from residual behavioural loss. This design is feasible for a three to four person team in eight weeks and remains informative if the proposed selector does not win.",
        style="Lead")

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()
